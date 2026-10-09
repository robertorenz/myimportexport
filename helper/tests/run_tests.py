"""
MIEConvert test suite.

Round-trips a tricky dataset through every format (export MIE-TSV -> format -> import MIE-TSV),
checks the generated xls/xlsx with independent readers (openpyxl, xlrd), reads files produced by
other writers (openpyxl, xlwt, a hand-built OpenXML package), and exercises the command-line
options (--replacesheet, sheets, --noheader, --maxrows, --rowpath, xml auto-detect, encodings).

Usage:
    python -m pip install openpyxl xlrd xlwt
    python run_tests.py [path\\to\\MIEConvert.exe]

Exit code 0 when every test passes.
"""
import datetime
import json
import os
import shutil
import subprocess
import sys
import tempfile
import traceback
import zipfile

HERE = os.path.dirname(os.path.abspath(__file__))
DEFAULT_EXE = os.path.normpath(os.path.join(HERE, "..", "..", "bin", "MIEConvert.exe"))
EXE = sys.argv[1] if len(sys.argv) > 1 else DEFAULT_EXE
WORK = tempfile.mkdtemp(prefix="mie_tests_")
ANSI = "cp1252"

try:
    import openpyxl
except ImportError:
    openpyxl = None
try:
    import xlrd
except ImportError:
    xlrd = None
try:
    import xlwt
except ImportError:
    xlwt = None

RESULTS = []


# ----------------------------------------------------------------------------- helpers

def p(name):
    return os.path.join(WORK, name)


def esc(s):
    return s.replace("\\", "\\\\").replace("\t", "\\t").replace("\n", "\\n").replace("\r", "\\r")


def unesc(s):
    out, i = [], 0
    while i < len(s):
        c = s[i]
        if c == "\\" and i + 1 < len(s) and s[i + 1] in "\\tnr":
            out.append({"\\": "\\", "t": "\t", "n": "\n", "r": "\r"}[s[i + 1]])
            i += 2
            continue
        out.append(c)
        i += 1
    return "".join(out)


def write_export_tsv(path, cols, types, rows):
    lines = ["\t".join(esc(c) for c in cols), "\t".join(types)]
    lines += ["\t".join(esc(v) for v in r) for r in rows]
    with open(path, "wb") as f:
        f.write(("\r\n".join(lines) + "\r\n").encode(ANSI))


def read_import_tsv(path):
    raw = open(path, "rb").read()
    assert b"\n" not in raw.replace(b"\r\n", b""), "bare LF in MIE-TSV output"
    text = raw.decode(ANSI)
    lines = text.split("\r\n")
    assert lines[-1] == "", "output must end with CRLF"
    lines = lines[:-1]
    cols = [unesc(x) for x in lines[0].split("\t")]
    rows = [[unesc(x) for x in ln.split("\t")] for ln in lines[1:]]
    return cols, rows


def run(*args, expect=0):
    errf = p("err.txt")
    if os.path.exists(errf):
        os.remove(errf)
    cmd = [EXE] + [str(a) for a in args] + ["--errfile", errf]
    r = subprocess.run(cmd, capture_output=True)
    if r.returncode != expect:
        msg = open(errf, encoding=ANSI).read() if os.path.exists(errf) else r.stderr.decode(ANSI, "replace")
        raise AssertionError("exit %d (expected %d) for %s\n%s" % (r.returncode, expect, " ".join(cmd[1:]), msg))
    return open(errf, encoding=ANSI).read() if os.path.exists(errf) else ""


def test(name):
    def deco(fn):
        def wrapper():
            try:
                detail = fn() or ""
                RESULTS.append((name, True, detail))
                print("PASS  %-55s %s" % (name, detail))
            except Exception as e:  # noqa
                RESULTS.append((name, False, str(e)))
                print("FAIL  %-55s %s" % (name, e))
                traceback.print_exc()
        wrapper.__name__ = fn.__name__
        TESTS.append(wrapper)
        return wrapper
    return deco


TESTS = []


def eq(a, b, what=""):
    if a != b:
        raise AssertionError("%s mismatch:\n  got      %r\n  expected %r" % (what, a, b))


def num_eq(a, b):
    if a == b:
        return True
    try:
        return abs(float(a) - float(b)) < 1e-9
    except ValueError:
        return False


def compare_tables(got_cols, got_rows, cols, types, rows, numeric_loose=False, label=""):
    eq(got_cols, cols, label + " columns")
    eq(len(got_rows), len(rows), label + " row count")
    for ri, (g, e) in enumerate(zip(got_rows, rows)):
        eq(len(g), len(e), "%s row %d width" % (label, ri))
        for ci, (gv, ev) in enumerate(zip(g, e)):
            if types[ci] == "N" and numeric_loose:
                if not num_eq(gv, ev) and not (gv == "" and ev.strip() == ""):
                    raise AssertionError("%s row %d col %s: got %r expected %r" % (label, ri, cols[ci], gv, ev))
            elif gv != ev:
                raise AssertionError("%s row %d col %s: got %r expected %r" % (label, ri, cols[ci], gv, ev))


# ----------------------------------------------------------------------------- sample data

COLS = ["ID", "Name", "Amount", "Born", "AtTime", "Notes", "Code", "Empty", "Línea ñ"]
TYPES = ["N", "S", "N", "D", "T", "S", "S", "S", "S"]
ROWS = [
    ["1", "Smith, John", "12.5", "1980-05-17", "08:30:00", 'He said "hi"', "00123", "", "a"],
    ["2", "José Muñoz", "-3", "2024-02-29", "23:59:59", "line1\nline2", "tab\there", "", "é ñ ü"],
    ["3", "Über ñandú", "0.125", "1900-03-01", "00:00:00", "back\\slash C:\\temp\\new", "  lead", "", "Ç"],
    ["4", "semi;colon|pipe", "1234567.89", "", "", "crlf\r\nend", "<xml> & 'stuff'", "", ""],
    ["5", "", "", "1999-12-31", "12:00:00", "x", "é", "notempty", "\\t literal"],
    ["6", "'quote", "100", "2000-01-01", "06:07:08", "=1+2", "true", "", "€ euro"],
    ["7", "trailing ", "0", "2038-01-19", "03:14:07", "a,b;c|d\te", "007", "", ",\"\n"],
]


def sample_tsv(name="sample.tsv"):
    path = p(name)
    write_export_tsv(path, COLS, TYPES, ROWS)
    return path


# ----------------------------------------------------------------------------- round trips

def expected_cols(fmt):
    # XML element names are sanitised on export, so the space in "Línea ñ" comes back as "_"
    return [c.replace(" ", "_") for c in COLS] if fmt == "xml" else COLS


def roundtrip(fmt, extra_export=(), extra_import=(), loose=False):
    src = sample_tsv()
    out = p("rt." + fmt)
    back = p("rt_%s_back.tsv" % fmt)
    run("export", src, out, *extra_export)
    run("import", out, back, *extra_import)
    cols, rows = read_import_tsv(back)
    compare_tables(cols, rows, expected_cols(fmt), TYPES, ROWS, numeric_loose=loose, label=fmt)
    return "%d rows x %d cols identical" % (len(rows), len(cols))


@test("round trip csv")
def t_csv():
    return roundtrip("csv")


@test("round trip csv ; delimiter + ansi encoding")
def t_csv_semi():
    src = sample_tsv()
    out = p("semi.csv")
    run("export", src, out, "--delimiter", ";", "--encoding", "ansi")
    raw = open(out, "rb").read()
    assert not raw.startswith(b"\xef\xbb\xbf"), "ansi output must not have a BOM"
    assert "José".encode(ANSI) in raw
    back = p("semi_back.tsv")
    run("import", out, back, "--delimiter", ";")  # auto-detect must fall back to ANSI
    cols, rows = read_import_tsv(back)
    compare_tables(cols, rows, COLS, TYPES, ROWS, label="csv;")
    return "ANSI file auto-detected, identical"


@test("round trip tsv")
def t_tsv():
    return roundtrip("tsv")


@test("round trip json")
def t_json():
    r = roundtrip("json")
    data = json.load(open(p("rt.json"), encoding="utf-8-sig"))
    eq(data[0]["Amount"], 12.5, "json number")
    eq(data[4]["Amount"], None, "json empty number -> null")
    eq(data[0]["ID"], 1, "json int")
    return r


@test("round trip json --root")
def t_json_root():
    src = sample_tsv()
    out = p("rt_root.json")
    run("export", src, out, "--root", "customers")
    data = json.load(open(out, encoding="utf-8-sig"))
    assert list(data.keys()) == ["customers"] and len(data["customers"]) == len(ROWS)
    back = p("rt_root_back.tsv")
    run("import", out, back, "--rowpath", "customers")
    cols, rows = read_import_tsv(back)
    compare_tables(cols, rows, COLS, TYPES, ROWS, label="json root")
    return "wrapped object, identical"


@test("round trip xml")
def t_xml():
    r = roundtrip("xml")
    import xml.etree.ElementTree as ET
    root = ET.parse(p("rt.xml")).getroot()
    eq(root.tag, "Data", "xml root")
    eq(root[0].tag, "Row", "xml row")
    eq(root[0][8].tag, "Línea_ñ", "sanitised element name")
    return r + "; ElementTree parses it"


@test("round trip xml --root/--row + sanitised names")
def t_xml_names():
    path = p("names.tsv")
    write_export_tsv(path, ["1st col", "a/b", "ok.name"], ["S", "S", "N"], [["x", "y", "1"]])
    out = p("names.xml")
    run("export", path, out, "--root", "My Root", "--row", "Item")
    import xml.etree.ElementTree as ET
    root = ET.parse(out).getroot()
    eq(root.tag, "My_Root", "root")
    eq([c.tag for c in root[0]], ["_1st_col", "a_b", "ok.name"], "names")
    return "My Root->My_Root, 1st col->_1st_col, a/b->a_b"


@test("round trip xlsx")
def t_xlsx():
    return roundtrip("xlsx", loose=True)


@test("round trip xls")
def t_xls():
    return roundtrip("xls", loose=True)


@test("xlsx numbers: 12.5 stays 12.5, integers have no .0")
def t_xlsx_numbers():
    cols, rows = read_import_tsv(p("rt_xlsx_back.tsv"))
    eq(rows[0][2], "12.5", "12.5")
    eq(rows[5][2], "100", "100")
    eq(rows[0][0], "1", "ID")
    eq(rows[3][2], "1234567.89", "1234567.89")
    eq(rows[2][2], "0.125", "0.125")
    return "12.5 / 100 / 1234567.89 / 0.125"


# ----------------------------------------------------------------------------- independent readers

@test("xlsx output opens in openpyxl with real dates/times/numbers")
def t_openpyxl_read():
    if openpyxl is None:
        raise AssertionError("openpyxl not installed")
    wb = openpyxl.load_workbook(p("rt.xlsx"))
    ws = wb.active
    eq(wb.sheetnames, ["Sheet1"], "sheet names")
    eq([c.value for c in ws[1]], COLS, "header")
    eq(ws["A2"].value, 1, "A2")
    eq(ws["C2"].value, 12.5, "C2")
    born = ws["D2"].value
    assert isinstance(born, datetime.datetime) and born.date() == datetime.date(1980, 5, 17), repr(born)
    eq(ws["D2"].number_format, "yyyy-mm-dd", "date format")
    t = ws["E2"].value
    tt = t if isinstance(t, datetime.time) else t.time()
    eq(tt, datetime.time(8, 30), "time")
    eq(ws["E2"].number_format, "hh:mm:ss", "time format")
    eq(ws["F3"].value, "line1\nline2", "newline")
    eq(ws["F5"].value, "crlf\r\nend", "crlf")
    eq(ws["G2"].value, "00123", "text that looks numeric")
    eq(ws["F7"].value, "=1+2", "formula-looking text stays text")
    eq(ws["G4"].value, "  lead", "leading spaces")
    assert ws["A1"].font.b, "header should be bold"
    return "dates are datetime, formats yyyy-mm-dd / hh:mm:ss, strings intact"


@test("xls output opens in xlrd with real dates/times/numbers")
def t_xlrd_read():
    if xlrd is None:
        raise AssertionError("xlrd not installed")
    book = xlrd.open_workbook(p("rt.xls"), formatting_info=True)
    eq(book.sheet_names(), ["Sheet1"], "sheet names")
    sh = book.sheet_by_index(0)
    eq(sh.row_values(0), COLS, "header")
    eq(sh.cell_value(1, 2), 12.5, "number")
    eq(sh.cell_type(1, 3), xlrd.XL_CELL_DATE, "date cell type")
    eq(xlrd.xldate_as_tuple(sh.cell_value(1, 3), book.datemode)[:3], (1980, 5, 17), "date value")
    eq(sh.cell_type(1, 4), xlrd.XL_CELL_DATE, "time cell type")
    eq(xlrd.xldate_as_tuple(sh.cell_value(1, 4), book.datemode)[3:], (8, 30, 0), "time value")
    eq(sh.cell_value(2, 1), "José Muñoz", "accents")
    eq(sh.cell_value(6, 8), "€ euro", "euro sign (UTF-16 string)")
    eq(sh.cell_value(4, 5), "crlf\r\nend", "crlf")
    xf = book.xf_list[sh.cell_xf_index(1, 3)]
    eq(book.format_map[xf.format_key].format_str, "yyyy-mm-dd", "date format string")
    font = book.font_list[book.xf_list[sh.cell_xf_index(0, 0)].font_index]
    assert font.bold, "header should be bold"
    return "xlrd: date cells XL_CELL_DATE, format yyyy-mm-dd, header bold"


def build_raw_xlsx(path):
    """Hand-built OpenXML package: shared strings (rich + phonetic), inline strings, sparse cells,
    built-in and custom date formats, booleans, formulas with cached values, a 1900 date system."""
    ct = ('<?xml version="1.0" encoding="UTF-8"?><Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
          '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
          '<Default Extension="xml" ContentType="application/xml"/>'
          '<Override PartName="/xl/workbook.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/>'
          '<Override PartName="/xl/worksheets/s1.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/>'
          '<Override PartName="/xl/worksheets/s2.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/>'
          '<Override PartName="/xl/sharedStrings.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sharedStrings+xml"/>'
          '<Override PartName="/xl/styles.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.styles+xml"/></Types>')
    rels = ('<?xml version="1.0" encoding="UTF-8"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="xl/workbook.xml"/></Relationships>')
    wb = ('<?xml version="1.0" encoding="UTF-8"?><workbook xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" '
          'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"><sheets>'
          '<sheet name="Info" sheetId="5" r:id="rIdB"/><sheet name="Datos ñ" sheetId="3" r:id="rIdA"/></sheets></workbook>')
    wbrels = ('<?xml version="1.0" encoding="UTF-8"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
              '<Relationship Id="rIdA" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" Target="/xl/worksheets/s1.xml"/>'
              '<Relationship Id="rIdB" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" Target="worksheets/s2.xml"/>'
              '<Relationship Id="rIdS" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/sharedStrings" Target="sharedStrings.xml"/>'
              '<Relationship Id="rIdT" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/styles" Target="styles.xml"/></Relationships>')
    sst = ('<?xml version="1.0" encoding="UTF-8"?><sst xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" count="6" uniqueCount="6">'
           '<si><t>Name</t></si><si><t>When</t></si><si><r><t>Ri</t></r><r><rPr><b/></rPr><t xml:space="preserve">ch </t></r><rPh><t>IGNORE</t></rPh></si>'
           '<si><t>Peña</t></si><si><t>Price</t></si><si><t>a_x000D_b</t></si></sst>')
    styles = ('<?xml version="1.0" encoding="UTF-8"?><styleSheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
              '<numFmts count="3"><numFmt numFmtId="170" formatCode="dd/mm/yyyy\\ hh:mm"/>'
              '<numFmt numFmtId="171" formatCode="&quot;Day&quot;\\ 0.00"/>'
              '<numFmt numFmtId="172" formatCode="[$-409]mmmm\\ d\\,\\ yyyy;@"/></numFmts>'
              '<fonts count="1"><font/></fonts><fills count="1"><fill/></fills><borders count="1"><border/></borders>'
              '<cellXfs count="7"><xf numFmtId="0"/><xf numFmtId="14"/><xf numFmtId="20"/><xf numFmtId="22"/>'
              '<xf numFmtId="170"/><xf numFmtId="171"/><xf numFmtId="172"/></cellXfs></styleSheet>')
    s1 = ('<?xml version="1.0" encoding="UTF-8"?><worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><sheetData>'
          '<row r="2"><c r="B2" t="s"><v>0</v></c><c r="C2" t="s"><v>1</v></c><c r="D2" t="s"><v>4</v></c><c r="E2" t="inlineStr"><is><t>Flag</t></is></c>'
          '<c r="F2" t="inlineStr"><is><t>Calc</t></is></c><c r="G2" t="inlineStr"><is><t>DT</t></is></c><c r="H2" t="inlineStr"><is><t>Custom</t></is></c>'
          '<c r="I2" t="inlineStr"><is><t>Txt</t></is></c><c r="J2" t="inlineStr"><is><t>LongDate</t></is></c></row>'
          '<row r="3"><c r="B3" t="s"><v>2</v></c><c r="C3" s="1"><v>45000</v></c><c r="D3"><v>12.5</v></c><c r="E3" t="b"><v>1</v></c>'
          '<c r="F3"><f>D3*2</f><v>25</v></c><c r="G3" s="3"><v>45000.75</v></c><c r="H3" s="5"><v>3</v></c><c r="I3" t="str"><f>"x"&amp;"y"</f><v>xy</v></c>'
          '<c r="J3" s="6"><v>45001</v></c></row>'
          '<row r="5"><c r="B5" t="s"><v>3</v></c><c r="C5" s="2"><v>0.5</v></c><c r="E5" t="b"><v>0</v></c><c r="G5" s="4"><v>36526.5</v></c>'
          '<c r="I5" t="s"><v>5</v></c></row>'
          '<row><c t="inlineStr"><is><t>noref</t></is></c><c><v>1E-3</v></c></row>'
          '</sheetData></worksheet>')
    s2 = ('<?xml version="1.0" encoding="UTF-8"?><worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><sheetData>'
          '<row r="1"><c r="A1" t="inlineStr"><is><t>k</t></is></c></row><row r="2"><c r="A2"><v>1</v></c></row></sheetData></worksheet>')
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("[Content_Types].xml", ct)
        z.writestr("_rels/.rels", rels)
        z.writestr("xl/workbook.xml", wb)
        z.writestr("xl/_rels/workbook.xml.rels", wbrels)
        z.writestr("xl/sharedStrings.xml", sst)
        z.writestr("xl/styles.xml", styles)
        z.writestr("xl/worksheets/s1.xml", s1)
        z.writestr("xl/worksheets/s2.xml", s2)


@test("import hand-built xlsx (sst, sparse, dates, bool, formulas)")
def t_raw_xlsx():
    path = p("raw.xlsx")
    build_raw_xlsx(path)
    out = p("raw_sheets.txt")
    run("sheets", path, out)
    eq(open(out, encoding=ANSI, newline="").read(), "Info\r\nDatos ñ\r\n", "sheets")
    back = p("raw_back.tsv")
    run("import", path, back, "--sheet", "datos ñ")
    cols, rows = read_import_tsv(back)
    eq(cols, ["Column1", "Name", "When", "Price", "Flag", "Calc", "DT", "Custom", "Txt", "LongDate"], "columns")
    eq(rows[0], ["", "Rich ", "2023-03-15", "12.5", "1", "25", "2023-03-15 18:00:00", "3", "xy", "2023-03-16"], "row 3")
    eq(rows[1], ["", "Peña", "12:00:00", "", "0", "", "2000-01-01 12:00:00", "", "a\rb", ""], "row 5 (sparse)")
    eq(rows[2], ["noref", "0.001", "", "", "", "", "", "", "", ""], "row without refs")
    return "header from first non-empty row, dates/times/datetime, bool 1/0, cached formulas, _x000D_"


@test("import openpyxl-written xlsx")
def t_openpyxl_written():
    if openpyxl is None:
        raise AssertionError("openpyxl not installed")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "People"
    ws.append(["Name", "Age", "Since", "Start", "Active", "Ratio"])
    ws.append(["Ana", 30, datetime.date(2020, 1, 31), datetime.time(9, 15, 30), True, 0.1])
    ws.append(["Bjørn", 41.5, datetime.datetime(1999, 12, 31, 23, 59, 59), datetime.time(0, 0, 1), False, 1e-7])
    ws["H5"] = "far"
    wb.create_sheet("Other")
    path = p("openpyxl.xlsx")
    wb.save(path)
    out = p("openpyxl_sheets.txt")
    run("sheets", path, out)
    eq(open(out, encoding=ANSI, newline="").read().split("\r\n")[:-1], ["People", "Other"], "sheets")
    back = p("openpyxl_back.tsv")
    run("import", path, back)
    cols, rows = read_import_tsv(back)
    eq(cols, ["Name", "Age", "Since", "Start", "Active", "Ratio", "Column7", "Column8"], "columns")
    eq(rows[0][:6], ["Ana", "30", "2020-01-31", "09:15:30", "1", "0.1"], "row 1")
    eq(rows[1][:6], ["Bjørn", "41.5", "1999-12-31 23:59:59", "00:00:01", "0", "0.0000001"], "row 2")
    eq(rows[2][7], "far", "sparse far cell")
    return "openpyxl dates/times/bools/sparse ok"


@test("import xlwt-written xls (SST/CONTINUE, RK, dates, bools)")
def t_xlwt_written():
    if xlwt is None:
        raise AssertionError("xlwt not installed")
    wb = xlwt.Workbook(encoding="utf-8")
    ws = wb.add_sheet("Hoja ñ")
    ws2 = wb.add_sheet("Second")
    dstyle = xlwt.easyxf(num_format_str="YYYY-MM-DD")
    tstyle = xlwt.easyxf(num_format_str="hh:mm:ss")
    dtstyle = xlwt.easyxf(num_format_str="dd/mm/yyyy hh:mm")
    hdr = ["Text", "Int", "Float", "Date", "Time", "DT", "Bool"]
    for c, h in enumerate(hdr):
        ws.write(0, c, h)
    long_text = "ü€" * 3000  # forces SST CONTINUE records with UTF-16 data
    expected = []
    for r in range(1, 301):
        txt = "row %d %s" % (r, "é" * (r % 7)) if r != 5 else long_text
        ws.write(r, 0, txt)
        ws.write(r, 1, r * 1000)
        ws.write(r, 2, r + 0.25)
        ws.write(r, 3, datetime.date(2001, 1, 1) + datetime.timedelta(days=r), dstyle)
        ws.write(r, 4, datetime.time(r % 24, r % 60, 0), tstyle)
        ws.write(r, 5, datetime.datetime(2010, 6, 1, 12, 30), dtstyle)
        ws.write(r, 6, bool(r % 2))
        expected.append([txt, str(r * 1000), repr(r + 0.25),
                         (datetime.date(2001, 1, 1) + datetime.timedelta(days=r)).isoformat(),
                         "%02d:%02d:00" % (r % 24, r % 60), "2010-06-01 12:30:00", "1" if r % 2 else "0"])
    ws2.write(0, 0, "x")
    path = p("xlwt.xls")
    wb.save(path)
    out = p("xlwt_sheets.txt")
    run("sheets", path, out)
    eq(open(out, encoding=ANSI, newline="").read().split("\r\n")[:-1], ["Hoja ñ", "Second"], "sheets")
    back = p("xlwt_back.tsv")
    run("import", path, back, "--codepage", "65001")  # UTF-8 MIE-TSV to keep the euro sign
    text = open(back, "rb").read().decode("utf-8")
    lines = text.split("\r\n")[:-1]
    cols = [unesc(x) for x in lines[0].split("\t")]
    rows = [[unesc(x) for x in ln.split("\t")] for ln in lines[1:]]
    eq(cols, hdr, "header")
    eq(len(rows), 300, "rows")
    for i, (g, e) in enumerate(zip(rows, expected)):
        eq(g, e, "row %d" % (i + 1))
    return "300 rows, 6000-char UTF-16 string across CONTINUE, dates/times/bools ok"


@test("large xls round trip (SST CONTINUE writer) + xlrd")
def t_big_xls():
    cols = ["N", "Text", "D"]
    types = ["N", "S", "D"]
    rows = []
    for i in range(3000):
        rows.append([str(i), ("Señal %d " % i) + ("x" * (i % 50)) + ("\n2nd" if i % 10 == 0 else ""), "2020-01-%02d" % (i % 28 + 1)])
    src = p("big.tsv")
    write_export_tsv(src, cols, types, rows)
    out = p("big.xls")
    run("export", src, out)
    back = p("big_back.tsv")
    run("import", out, back)
    gc, gr = read_import_tsv(back)
    compare_tables(gc, gr, cols, types, rows, numeric_loose=True, label="big xls")
    if xlrd is not None:
        sh = xlrd.open_workbook(out).sheet_by_index(0)
        eq(sh.nrows, 3001, "xlrd rows")
        eq(sh.cell_value(2999, 1), rows[2998][1], "xlrd string")
    return "3000 rows, %d KB, xlrd agrees" % (os.path.getsize(out) // 1024)


@test("large xlsx round trip")
def t_big_xlsx():
    cols, types, rows = ["N", "Text"], ["N", "S"], [[str(i * 1.5), "r%d é" % i] for i in range(20000)]
    src = p("bigx.tsv")
    write_export_tsv(src, cols, types, rows)
    out = p("big.xlsx")
    run("export", src, out)
    back = p("bigx_back.tsv")
    run("import", out, back)
    gc, gr = read_import_tsv(back)
    compare_tables(gc, gr, cols, types, rows, numeric_loose=True, label="big xlsx")
    return "20000 rows"


# ----------------------------------------------------------------------------- options

@test("--replacesheet keeps other sheets (xlsx)")
def t_replacesheet():
    out = p("multi.xlsx")
    a1, b1, a2 = p("a1.tsv"), p("b1.tsv"), p("a2.tsv")
    write_export_tsv(a1, ["X"], ["S"], [["first"]])
    write_export_tsv(b1, ["Y", "When"], ["N", "D"], [["42", "2021-07-04"]])
    write_export_tsv(a2, ["X", "Z"], ["S", "T"], [["second", "10:11:12"], ["third", ""]])
    run("export", a1, out, "--sheet", "Alpha")
    run("export", b1, out, "--sheet", "Beta", "--replacesheet")
    run("export", a2, out, "--sheet", "Alpha", "--replacesheet")
    sl = p("multi_sheets.txt")
    run("sheets", out, sl)
    eq(open(sl, encoding=ANSI, newline="").read(), "Alpha\r\nBeta\r\n", "sheets")
    back = p("multi_a.tsv")
    run("import", out, back, "--sheet", "Alpha")
    eq(read_import_tsv(back), (["X", "Z"], [["second", "10:11:12"], ["third", ""]]), "Alpha replaced")
    run("import", out, back, "--sheet", "Beta")
    eq(read_import_tsv(back), (["Y", "When"], [["42", "2021-07-04"]]), "Beta kept")
    if openpyxl is not None:
        wb = openpyxl.load_workbook(out)
        eq(wb.sheetnames, ["Alpha", "Beta"], "openpyxl sheets")
        eq(wb["Beta"]["B2"].number_format, "yyyy-mm-dd", "Beta date format kept")
        eq(wb["Alpha"]["B2"].number_format, "hh:mm:ss", "Alpha time format")
    # without --replacesheet the file is overwritten
    run("export", b1, out, "--sheet", "Only")
    run("sheets", out, sl)
    eq(open(sl, encoding=ANSI, newline="").read(), "Only\r\n", "overwrite")
    return "Alpha replaced, Beta intact, openpyxl agrees; plain export overwrites"


@test("--replacesheet into a workbook made by openpyxl")
def t_replacesheet_foreign():
    if openpyxl is None:
        raise AssertionError("openpyxl not installed")
    from openpyxl.styles import Font, PatternFill
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Keep"
    ws["A1"] = "styled"
    ws["A1"].font = Font(italic=True, color="FF0000")
    ws["A1"].fill = PatternFill("solid", fgColor="FFFF00")
    ws["B1"] = "=1+1"
    ws["C1"] = datetime.date(2022, 2, 2)
    wb.create_sheet("Data")["A1"] = "old"
    path = p("foreign.xlsx")
    wb.save(path)
    src = sample_tsv()
    run("export", src, path, "--sheet", "Data", "--replacesheet")
    run("export", src, path, "--sheet", "New Sheet", "--replacesheet")
    wb2 = openpyxl.load_workbook(path)
    eq(wb2.sheetnames, ["Keep", "Data", "New Sheet"], "sheets")
    k = wb2["Keep"]
    eq(k["A1"].value, "styled", "kept value")
    assert k["A1"].font.i and k["A1"].fill.fgColor.rgb.endswith("FFFF00"), "kept style"
    eq(k["B1"].value, "=1+1", "kept formula")
    eq(wb2["Data"]["A1"].value, "ID", "replaced data")
    eq(wb2["Data"]["D2"].number_format, "yyyy-mm-dd", "merged date style")
    back = p("foreign_back.tsv")
    run("import", path, back, "--sheet", "New Sheet")
    cols, rows = read_import_tsv(back)
    compare_tables(cols, rows, COLS, TYPES, ROWS, numeric_loose=True, label="foreign")
    return "styles merged, foreign sheet + formula kept"


@test("sheets command (xls) + unknown sheet error")
def t_sheets_xls():
    src = sample_tsv()
    out = p("named.xls")
    run("export", src, out, "--sheet", "Clientes ñ")
    sl = p("named_sheets.txt")
    run("sheets", out, sl)
    eq(open(sl, encoding=ANSI, newline="").read(), "Clientes ñ\r\n", "sheets")
    msg = run("import", out, p("x.tsv"), "--sheet", "Nope", expect=1)
    assert "Nope" in msg, msg
    msg = run("sheets", p("rt.csv"), sl, expect=1)
    assert "xls" in msg, msg
    return "sheet name with ñ; errors reported to --errfile"


@test("--noheader import and export")
def t_noheader():
    src = sample_tsv()
    out = p("nohdr.csv")
    run("export", src, out, "--noheader")
    first = open(out, encoding="utf-8-sig").readline()
    assert first.startswith("1,"), first
    back = p("nohdr_back.tsv")
    run("import", out, back, "--noheader")
    cols, rows = read_import_tsv(back)
    eq(cols, ["Column%d" % i for i in range(1, 10)], "columns")
    compare_tables(cols, rows, cols, TYPES, ROWS, label="noheader")
    for fmt in ("xlsx", "xls"):
        o = p("nohdr." + fmt)
        run("export", src, o, "--noheader")
        run("import", o, back, "--noheader")
        c2, r2 = read_import_tsv(back)
        eq(len(r2), len(ROWS), fmt + " rows")
        eq(r2[1][1], "José Muñoz", fmt + " value")
        eq(c2[0], "Column1", fmt + " col")
    return "csv/xlsx/xls"


@test("--maxrows on every format")
def t_maxrows():
    out = []
    for fmt in ("csv", "tsv", "json", "xml", "xlsx", "xls"):
        back = p("max_%s.tsv" % fmt)
        run("import", p("rt." + fmt), back, "--maxrows", "2")
        cols, rows = read_import_tsv(back)
        eq(len(rows), 2, fmt)
        eq(cols, expected_cols(fmt), fmt + " columns")
        out.append(fmt)
    return "2 rows from " + "/".join(out)


@test("json: nested objects flattened, arrays as JSON text, --rowpath")
def t_json_nested():
    doc = {
        "meta": {"count": 2, "src": "test"},
        "data": {"items": [
            {"id": 1, "addr": {"city": "Zürich", "geo": {"lat": 47.37, "lon": 8.54}}, "tags": ["a", "b"], "ok": True, "n": None, "big": 1.5e20},
            {"id": 2, "extra": "y\"q", "ok": False},
        ]},
        "other": [{"z": 1}],
    }
    path = p("nested.json")
    with open(path, "w", encoding="utf-8") as f:
        json.dump(doc, f, ensure_ascii=False)
    back = p("nested_back.tsv")
    run("import", path, back, "--rowpath", "items")
    cols, rows = read_import_tsv(back)
    eq(cols, ["id", "addr.city", "addr.geo.lat", "addr.geo.lon", "tags", "ok", "n", "big", "extra"], "columns")
    eq(rows[0], ["1", "Zürich", "47.37", "8.54", '["a","b"]', "1", "", "150000000000000000000", ""], "row 1")
    eq(rows[1], ["2", "", "", "", "", "0", "", "", 'y"q'], "row 2")
    run("import", path, back)  # no rowpath: first array found (data.items)
    eq(read_import_tsv(back)[0][0], "id", "auto first array")
    run("import", path, back, "--rowpath", "data.items")
    eq(read_import_tsv(back)[0][0], "id", "dotted path")
    run("import", path, back, "--rowpath", "other")
    eq(read_import_tsv(back), (["z"], [["1"]]), "other array")
    msg = run("import", path, back, "--rowpath", "missing", expect=1)
    assert "missing" in msg
    return "parent.child, arrays as JSON, bool 1/0, null empty"


@test("json: root array of scalars / arrays, BOM, utf-16")
def t_json_misc():
    path = p("arr.json")
    open(path, "wb").write('[["a","b"],[1,"ü"],[2,null]]'.encode("utf-16"))
    back = p("arr_back.tsv")
    run("import", path, back)
    eq(read_import_tsv(back), (["a", "b"], [["1", "ü"], ["2", ""]]), "array of arrays utf-16")
    open(path, "wb").write(b"\xef\xbb\xbf[1,2,3]")
    run("import", path, back)
    eq(read_import_tsv(back), (["Value"], [["1"], ["2"], ["3"]]), "scalars")
    return "ok"


@test("xml: auto-detect repeating rows, attributes + child text")
def t_xml_auto():
    path = p("auto.xml")
    open(path, "w", encoding="utf-8").write(
        '<?xml version="1.0" encoding="utf-8"?>\n'
        '<Export generated="today"><Header><Title>Customers</Title><By>me</By></Header>\n'
        '<Customers>\n'
        '  <Customer id="1" vip="yes"><Name>Ann &amp; Co</Name><Addr><City>Köln</City><Zip>50667</Zip></Addr></Customer>\n'
        '  <Customer id="2"><Name>Bob</Name><Note>  spaced  </Note></Customer>\n'
        '  <Customer id="3"><Name/></Customer>\n'
        '</Customers></Export>')
    back = p("auto_back.tsv")
    run("import", path, back)
    cols, rows = read_import_tsv(back)
    eq(cols, ["id", "vip", "Name", "Addr.City", "Addr.Zip", "Note"], "columns")
    eq(rows, [["1", "yes", "Ann & Co", "Köln", "50667", ""],
              ["2", "", "Bob", "", "", "  spaced  "],
              ["3", "", "", "", "", ""]], "rows")
    path2 = p("leaf.xml")
    open(path2, "w", encoding="utf-8").write("<list><item>a</item><item>b</item><item>c</item></list>")
    run("import", path2, back)
    eq(read_import_tsv(back), (["item"], [["a"], ["b"], ["c"]]), "leaf rows")
    run("import", path, back, "--rowpath", "Header")
    eq(read_import_tsv(back), (["Title", "By"], [["Customers", "me"]]), "--rowpath Header")
    return "Customer detected, attributes+children, nested flattened"


@test("csv: RFC 4180 + encodings (utf-8, utf-8 BOM, ansi, utf-16)")
def t_csv_rfc():
    text = 'a,"b ""q""",c\r\n"multi\r\nline","x,y",\r\n\r\n"",é,"ñ"\n'
    expected = (["a", 'b "q"', "c"], [["multi\r\nline", "x,y", ""], ["", "é", "ñ"]])
    back = p("rfc_back.tsv")
    for enc, data in (("utf8", text.encode("utf-8")), ("utf8bom", b"\xef\xbb\xbf" + text.encode("utf-8")),
                      ("ansi", text.encode(ANSI)), ("utf16", text.encode("utf-16"))):
        path = p("rfc_%s.csv" % enc)
        open(path, "wb").write(data)
        run("import", path, back)
        eq(read_import_tsv(back), expected, enc)
    path = p("rfc_utf8.csv")
    run("import", path, back, "--encoding", "1252")
    assert read_import_tsv(back)[1][1][1] != "é", "forced 1252 should mis-decode utf-8"
    path = p("pipe.txt")
    open(path, "w", encoding="utf-8").write("h1|h2\n1|2\n")
    run("import", path, back, "--format", "csv", "--delimiter", "|")
    eq(read_import_tsv(back), (["h1", "h2"], [["1", "2"]]), "pipe")
    return "quotes, doubled quotes, embedded CRLF, blank line skipped, 4 encodings"


@test("export --encoding utf8 / ansi / utf8bom")
def t_out_enc():
    src = sample_tsv()
    for enc, bom, codec in (("utf8", False, "utf-8"), ("utf8bom", True, "utf-8"), ("ansi", False, ANSI), ("1252", False, ANSI)):
        out = p("enc_%s.tsv.csv" % enc)
        run("export", src, out, "--encoding", enc)
        raw = open(out, "rb").read()
        eq(raw.startswith(b"\xef\xbb\xbf"), bom, enc + " bom")
        assert "José Muñoz" in raw.decode(codec, "replace"), enc
    return "ok"


@test("errors: missing file, bad option, exit code 1 + errfile")
def t_errors():
    msg = run("import", p("does_not_exist.csv"), p("x.tsv"), expect=1)
    assert "not found" in msg.lower(), msg
    msg = run("import", p("rt.csv"), p("x.tsv"), "--bogus", "1", expect=1)
    assert "bogus" in msg, msg
    msg = run("frobnicate", p("rt.csv"), p("x.tsv"), expect=1)
    assert "frobnicate" in msg, msg
    bad = p("bad.json")
    open(bad, "w").write('{"a": [1, 2')
    msg = run("import", bad, p("x.tsv"), expect=1)
    assert "JSON" in msg, msg
    return "messages written to errfile"


@test("empty dataset exports and imports")
def t_empty():
    src = p("empty_src.tsv")
    write_export_tsv(src, ["A", "B"], ["S", "N"], [])
    for fmt in ("csv", "tsv", "json", "xml", "xlsx", "xls"):
        out = p("empty." + fmt)
        run("export", src, out)
        back = p("empty_back.tsv")
        run("import", out, back)
        cols, rows = read_import_tsv(back)
        eq(rows, [], fmt)
        if fmt in ("csv", "tsv", "xlsx", "xls"):
            eq(cols, ["A", "B"], fmt + " cols")
    return "all formats"


# ----------------------------------------------------------------------------- main

def main():
    if not os.path.exists(EXE):
        print("MIEConvert.exe not found: " + EXE)
        return 2
    print("Testing " + EXE)
    print("Work dir " + WORK)
    for t in TESTS:
        t()
    passed = sum(1 for r in RESULTS if r[1])
    print("\n%d/%d tests passed" % (passed, len(RESULTS)))
    if passed == len(RESULTS):
        shutil.rmtree(WORK, ignore_errors=True)
    return 0 if passed == len(RESULTS) else 1


if __name__ == "__main__":
    sys.exit(main())
